import io
import json

import pytest
from fastapi.testclient import TestClient

from hack import workflow as W
from hack.api import create_app
from hack.challenge import contracts as C
from hack.config import Settings
from hack.llm.fake import FakeProvider, call_tools, text
from test_contract_review import GOV, LIAB, scripted_review


def settings(tmp_path, **kw) -> Settings:
    return Settings(
        provider="fake",
        runs_dir=tmp_path / "runs",
        web_dir=tmp_path / "web",
        uploads_dir=tmp_path / "uploads",
        results_dir=tmp_path / "results",
        proof_dir=tmp_path / "proof",
        **kw,
    )


def scripted_escalation(doc="nda-en-01"):
    """Maps every clause as Other: no liability/governing-law clause declared -> escalate."""
    ids = [f"c{i}" for i in range(1, 8)]
    return [
        call_tools(("1", "load_contract", {"doc_id": doc})),
        call_tools(
            (
                "2",
                "submit_clause_map",
                {
                    "doc_id": doc,
                    "clauses": [{"clause_id": i, "clause_types": ["Other"]} for i in ids],
                },
            )
        ),
        call_tools(("3", "score_contract", {"doc_id": doc})),
        call_tools(
            ("4", "escalate_to_legal", {"doc_id": doc, "question": "Which clause caps liability?"})
        ),
        text("ROUTE: escalate"),
    ]


# --- upload / ingest ---------------------------------------------------------------------------


def test_pdf_style_text_is_regrouped_into_clauses():
    raw = "SERVICE AGREEMENT\n1. Services. The supplier\nshall provide services.\n2. Fees. Net 30.\n3. Law. Saudi law."
    md = C.to_contract_markdown(raw, "hint")
    title, clauses = C.split_clauses(md)
    assert title == "SERVICE AGREEMENT"
    assert [c.text for c in clauses] == [
        "1. Services. The supplier shall provide services.",
        "2. Fees. Net 30.",
        "3. Law. Saudi law.",
    ]
    ar = "عقد خدمات\nالمادة 1: الخدمات يقدمها المورد.\nالمادة 2: يخضع العقد لأنظمة المملكة."
    assert len(C.split_clauses(C.to_contract_markdown(ar, "x"))[1]) == 2


def test_extract_text_docx_and_rejections():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Supply Agreement")
    doc.add_paragraph("1. Supply. The supplier delivers goods monthly.")
    doc.add_paragraph("2. Liability. Each party's liability is capped at fees paid in 12 months.")
    buf = io.BytesIO()
    doc.save(buf)
    raw = C.extract_text("supply.docx", buf.getvalue())
    title, clauses = C.split_clauses(C.to_contract_markdown(raw, "x"))
    assert title == "Supply Agreement" and len(clauses) == 2
    with pytest.raises(ValueError, match="unsupported"):
        C.extract_text("x.exe", b"MZ")
    with pytest.raises(ValueError, match="empty"):
        C.to_contract_markdown("   ", "x")
    with pytest.raises(ValueError, match="too large"):
        C.extract_text("x.txt", b"a" * (C.MAX_UPLOAD_BYTES + 1))
    assert C.extract_text("x.txt", "عقد".encode("cp1256")) == "عقد"  # legacy Arabic encoding


def test_runs_can_be_listed_for_one_contract(tmp_path):
    """Picking a contract opens its latest review: the API lists one contract's runs, newest first."""
    s = settings(tmp_path)
    client = TestClient(create_app(s))
    run_batch(s, lambda: scripted_review("services-en-01"), ["services-en-01"], False)
    run_batch(s, scripted_escalation, ["nda-en-01"], False)
    mine = client.get("/runs", params={"doc_id": "services-en-01"}).json()
    assert [r["doc_id"] for r in mine] == ["services-en-01"] and mine[0]["status"] == "done"
    assert client.get("/runs", params={"doc_id": "never-reviewed"}).json() == []
    assert client.get("/runs", params={"doc_id": "../etc"}).status_code == 400


def test_upload_endpoint_then_review_the_uploaded_contract(tmp_path):
    s = settings(tmp_path)
    client = TestClient(create_app(s))
    pasted = (
        "My NDA\n\n1. Confidentiality lasts four years.\n\n2. Saudi law governs this agreement."
    )
    r = client.post("/contracts", data={"text": pasted})
    assert r.status_code == 200, r.text
    doc = r.json()
    assert doc["source"] == "upload" and doc["clauses"] == 2 and doc["title"] == "My NDA"
    assert (s.uploads_dir / f"{doc['doc_id']}.md").is_file()
    listed = {d["doc_id"]: d["source"] for d in client.get("/contracts").json()}
    assert listed[doc["doc_id"]] == "upload" and listed["nda-en-01"] == "sample"
    assert client.get(f"/contracts/{doc['doc_id']}").json()["clauses"][1]["id"] == "c2"
    assert client.post("/contracts", data={"text": "  "}).status_code == 400
    bad = client.post("/contracts", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert bad.status_code == 400
    # the agent's tools find uploaded contracts too
    from hack.challenge.tools import build_registry
    from hack.llm.base import ToolCall

    res = build_registry(s).execute(ToolCall("x", "load_contract", {"doc_id": doc["doc_id"]}))
    assert not res.is_error and json.loads(res.content)["clause_count"] == 2


# --- batch, inbox, decisions -------------------------------------------------------------------


def run_batch(s, script_fn, docs, auto_send):
    if len(set(docs)) != len(docs):
        # The same contract twice cannot be in one batch (one review at a time per contract), so
        # run it as batches one after the other and return their items together.
        items, runner = [], None
        for d in docs:
            runner, st = run_batch(s, script_fn, [d], auto_send)
            items += st["items"]
        return runner, {"items": items}
    runner = W.BatchRunner(s, provider_factory=lambda: FakeProvider(script_fn()), workers=2)
    manifest = runner.start(docs, auto_send=auto_send)
    runner.pool.shutdown(wait=True)
    return runner, runner.status(manifest["batch_id"])


def test_batch_with_human_approval_of_sends(tmp_path):
    s = settings(tmp_path)
    # The same contract twice, one batch after the other (one batch may not hold it twice).
    for _ in range(2):
        _, st = run_batch(s, lambda: scripted_review("services-en-01"), ["services-en-01"], False)
        assert st["done"] == 1 and st["routes"]["negotiate"] == 1 and st["errors"] == 0
        assert st["needs_human"] == 1 and st["items"][0]["pending"] == "approve_send"
        assert st["estimate"]["hours_saved"] == 1.0 and "60 min" in st["estimate"]["assumptions"]
    items = W.inbox(s)
    assert {i["kind"] for i in items} == {"approve_send"} and len(items) == 2
    rid = items[0]["run_id"]
    with pytest.raises(ValueError, match="must be one of"):
        W.decide(s, rid, "approve_contract")
    rec = W.decide(s, rid, "send", note="ok")
    # Built from the recorded redline, one numbered line per change; not the model's "amend c3".
    assert rec["message"].startswith("Requested changes to ")
    assert "\n1. Clause 3 (Termination): " in rec["message"]
    out = [json.loads(x) for x in (s.runs_dir / rid / "outbox.jsonl").read_text().splitlines()]
    assert out[-1]["by"] == "human" and out[-1]["kind"] == "counterparty"
    with pytest.raises(ValueError, match="already decided"):
        W.decide(s, rid, "send")
    assert len(W.inbox(s)) == 1


def test_one_review_at_a_time_per_contract():
    """Two people (or a person and a batch) cannot review the same contract at once; the second
    is told who has it. Releasing takes the right run's hold only."""
    locks = W.ReviewLocks()
    locks.acquire("services-en-01", "mohammad", "r1")
    with pytest.raises(W.ReviewBusy, match="being reviewed by mohammad"):
        locks.acquire("services-en-01", "sara", "r2")
    locks.acquire("nda-en-01", "sara", "r3")  # another contract is free
    locks.release("services-en-01", "r2")  # not the holder: no effect
    assert locks.holder("services-en-01")["by"] == "mohammad"
    locks.release("services-en-01", "r1")
    assert locks.holder("services-en-01") is None
    locks.acquire("services-en-01", "sara", "r4")


def test_batch_refuses_a_contract_someone_is_reviewing_and_starts_nothing(tmp_path):
    locks = W.ReviewLocks()
    runner = W.BatchRunner(settings(tmp_path), locks=locks)
    locks.acquire("services-en-01", "mohammad", "live-1")
    with pytest.raises(W.ReviewBusy, match="mohammad"):
        runner.start(["nda-en-01", "services-en-01"], False, by="sara")
    assert locks.holder("nda-en-01") is None  # all or nothing: the free one was not kept
    with pytest.raises(ValueError, match="twice"):
        runner.start(["nda-en-01", "nda-en-01"], False)


def test_api_refuses_a_second_review_and_a_relabel_while_one_runs(tmp_path):
    """Through the API: while a review holds a contract, another review of it is refused with who
    has it, its label cannot change, and the contract list shows who is reviewing it."""
    import threading

    started, release = threading.Event(), threading.Event()

    class Slow(FakeProvider):
        def complete(self, *a, **kw):
            started.set()
            release.wait(5)
            return super().complete(*a, **kw)

    s = settings(tmp_path)
    client = TestClient(
        create_app(s, provider_factory=lambda: Slow(scripted_review("services-en-01")))
    )
    up = client.post("/contracts", data={"text": "1. Fees. Net 60.\n\n2. Law. Saudi law."}).json()
    doc = up["doc_id"]
    first = threading.Thread(
        target=lambda: client.post(
            "/run", json={"task": f"Review contract {doc} and route it.", "run_id": "first-1"}
        )
    )
    first.start()
    assert started.wait(5)
    second = client.post(
        "/run", json={"task": f"Review contract {doc} and route it.", "run_id": "second-1"}
    )
    assert second.status_code == 409 and "being reviewed" in second.json()["detail"]
    listed = {c["doc_id"]: c for c in client.get("/contracts").json()}
    assert listed[doc]["in_review"]["run_id"] == "first-1"
    relabel = client.put(f"/contracts/{doc}/sensitivity", json={"sensitivity": "restricted"})
    assert relabel.status_code == 409
    release.set()
    first.join(10)
    listed = {c["doc_id"]: c for c in client.get("/contracts").json()}
    assert listed[doc]["in_review"] is None  # released when the run ends


def test_batch_auto_send_needs_no_human_and_escalations_reach_legal(tmp_path):
    s = settings(tmp_path)
    _, st = run_batch(s, lambda: scripted_review("services-en-01"), ["services-en-01"], True)
    assert st["needs_human"] == 0 and st["items"][0]["sent"] is True
    _, st = run_batch(s, scripted_escalation, ["nda-en-01"], True)
    assert st["routes"]["escalate"] == 1 and st["items"][0]["pending"] == "legal_review"
    assert st["estimate"]["human_hours_with_agent"] == 0.25
    legal = [i for i in W.inbox(s) if i["kind"] == "legal_review"]
    assert len(legal) == 1 and "missing" in legal[0]["reasons"][0]
    rec = W.decide(s, legal[0]["run_id"], "request_changes")
    assert rec["message"].startswith("Requested changes to Mutual Non-Disclosure Agreement")


def test_batch_rejects_unknown_or_unsafe_docs(tmp_path):
    runner = W.BatchRunner(settings(tmp_path))
    with pytest.raises(FileNotFoundError):
        runner.start(["nope"], True)
    with pytest.raises(ValueError):
        runner.start(["../x"], True)
    with pytest.raises(ValueError):
        runner.start([], True)
    with pytest.raises(FileNotFoundError):
        runner.status("../../etc")


def test_failed_run_is_closed_not_running_forever(tmp_path):
    s = settings(tmp_path)

    def boom(_messages):
        raise RuntimeError("provider down")

    runner, st = run_batch(s, lambda: [boom], ["nda-en-01"], True)
    assert st["errors"] == 1 and st["items"][0]["status"] == "error"


def test_eval_results_pick_the_most_complete_run_per_model(tmp_path):
    s = settings(tmp_path)
    s.results_dir.mkdir()

    def write(name, model, ids):
        cases = [
            {"id": i, "passed": True, "expected_route": "approve", "got_route": "approve"}
            for i in ids
        ]
        summary = {"model": model, "route_accuracy": 1.0, "cases": len(ids)}
        (s.results_dir / name).write_text(json.dumps({"summary": summary, "cases": cases}))

    write("a.json", "m1", ["nda-en-01"])
    write("b.json", "m1", ["nda-en-01", "saas-en-01"])
    write("c.json", "m2", ["calc-1"])  # not the contract suite: ignored
    res = W.eval_results(s)
    assert [(r["model"], r["file"]) for r in res] == [("m1", "b.json")]


def test_api_workflow_endpoints(tmp_path):
    s = settings(tmp_path)
    client = TestClient(create_app(s, provider_factory=lambda: FakeProvider(scripted_escalation())))
    b = client.post("/batches", json={"doc_ids": ["nda-en-01"], "auto_send": False})
    assert b.status_code == 200
    assert client.post("/batches", json={"doc_ids": ["nope"]}).status_code == 404
    assert client.get("/batches/zzz").status_code == 404
    assert client.get("/evals").json() == []
    assert client.get("/runs/x/review").status_code == 404
    assert client.post("/runs/nope/decision", json={"decision": "send"}).status_code == 404
    assert client.get("/healthz").json()["approve_writes"] is False
    assert LIAB and GOV  # shared constants imported for parity with the contract tests


def test_junk_text_is_not_a_contract_and_eval_runs_never_pend(tmp_path):
    with pytest.raises(ValueError, match="fewer than 2 clauses"):
        C.to_contract_markdown("Title\n\n---\n\n... ---", "x")
    s = settings(tmp_path)
    runner, st = run_batch(s, lambda: scripted_review("services-en-01"), ["services-en-01"], False)
    assert st["priced"] is False  # the fake model has no price
    rid = st["items"][0]["run_id"]
    assert W.pending_kind(W.run_state(s, rid), None) == "approve_send"
    eval_state = {**W.run_state(s, rid), "run_id": "eval-x-1"}
    assert W.pending_kind(eval_state, None) is None


def test_static_catch_all_never_serves_files_outside_web_root(tmp_path):
    s = settings(tmp_path)
    s.web_dir.mkdir(parents=True)
    (s.web_dir / "index.html").write_text("<html>app</html>")
    (s.web_dir / "favicon.svg").write_text("<svg/>")
    (tmp_path / "secret.txt").write_text("TOP SECRET")
    client = TestClient(create_app(s))
    assert client.get("/favicon.svg").text == "<svg/>"
    for p in (
        "/..%2Fsecret.txt",
        "/../secret.txt",
        "/%2e%2e/secret.txt",
        "/assets/..%2F..%2Fsecret.txt",
    ):
        r = client.get(p)
        assert "TOP SECRET" not in r.text, p
    assert client.get("/some/spa/route").text == "<html>app</html>"


def test_missing_asset_is_404_not_html(tmp_path):
    s = settings(tmp_path)
    s.web_dir.mkdir(parents=True)
    (s.web_dir / "index.html").write_text("<html>app</html>")
    client = TestClient(create_app(s))
    assert client.get("/assets/old-build.js").status_code == 404


def test_corrupt_uploads_are_400_and_invalid_decisions_400(tmp_path):
    with pytest.raises(ValueError, match="could not be read"):
        C.extract_text("x.docx", b"not a zip")
    with pytest.raises(ValueError, match="could not be read"):
        C.extract_text("x.pdf", b"%PDF-1.4 garbage")
    s = settings(tmp_path)
    client = TestClient(create_app(s))
    r = client.post("/contracts", files={"file": ("x.pdf", b"%PDF-1.4 garbage", "application/pdf")})
    assert r.status_code == 400 and "could not be read" in r.json()["detail"]
    _, st = run_batch(s, lambda: scripted_review("services-en-01"), ["services-en-01"], False)
    rid = st["items"][0]["run_id"]
    assert client.post(f"/runs/{rid}/decision", json={"decision": "nonsense"}).status_code == 400
    assert client.post(f"/runs/{rid}/decision", json={"decision": "send"}).status_code == 200
    assert client.post(f"/runs/{rid}/decision", json={"decision": "send"}).status_code == 409


def test_polling_a_run_that_has_not_started_yet_is_not_404(tmp_path):
    import threading

    s = settings(tmp_path)
    gate = threading.Event()

    def slow(_messages):
        gate.wait(5)
        return text("done")

    app = create_app(s, provider_factory=lambda: FakeProvider([slow]))
    client = TestClient(app)
    t = threading.Thread(
        target=lambda: client.post("/run", json={"task": "x", "run_id": "ui-slow-1"})
    )
    t.start()
    try:
        import time as _t

        for _ in range(50):  # until the request is in flight
            r = client.get("/runs/ui-slow-1/review")
            if r.status_code == 200:
                break
            _t.sleep(0.05)
        assert r.status_code == 200 and r.json()["status"] in ("starting", "running")
        assert client.get("/runs/ui-slow-1/trace").status_code == 200
        assert client.get("/runs/never-started/review").status_code == 404
    finally:
        gate.set()
        t.join()


def test_proof_prefers_curated_files_and_keeps_variants_apart(tmp_path):
    s = settings(tmp_path)
    for d in (s.results_dir, s.proof_dir):
        d.mkdir()

    def write(folder, name, model, variant=None, acc=1.0):
        summary = {"model": model, "route_accuracy": acc, "cases": 1}
        if variant:
            summary["variant"] = variant
        cases = [{"id": "nda-en-01", "passed": True}]
        (folder / name).write_text(json.dumps({"summary": summary, "cases": cases}))

    write(s.results_dir, "local.json", "scratch-model")
    assert [r["model"] for r in W.eval_results(s)] == ["scratch-model"]  # no proof yet: fallback
    write(s.proof_dir, "a.json", "m", acc=0.75)
    write(s.proof_dir, "b.json", "m", variant="before red-flag tripwires", acc=0.58)
    res = W.eval_results(s)
    assert [(r["model"], r["variant"]) for r in res] == [
        ("m", None),
        ("m", "before red-flag tripwires"),
    ]


def stops_after_scoring(doc="services-en-01"):
    """Scores to negotiate, then quits without render/send; a second stop after the nudge."""
    full = scripted_review(doc)
    return [*full[:5], text("done, I think"), text("still done")]


def test_model_that_stops_early_is_nudged_to_finish(tmp_path):
    s = settings(tmp_path)
    script = [
        *scripted_review("services-en-01")[:5],  # up to and including score_contract
        text("done, I think"),
        *scripted_review("services-en-01")[5:],
    ]
    from hack.service import run_task

    r = run_task("Review contract services-en-01 and route it.", s, provider=FakeProvider(script))
    tr = [json.loads(x) for x in (s.runs_dir / r.run_id / "trace.jsonl").read_text().splitlines()]
    nudges = [e for e in tr if e["type"] == "nudge"]
    assert len(nudges) == 1 and "send_to_counterparty" in nudges[0]["text"]
    assert "send_to_counterparty" in r.tool_names  # it finished after the nudge


def test_nudges_are_bounded_and_the_workflow_owns_the_follow_up(tmp_path):
    s = settings(tmp_path)
    # human mode: the model never sends -> the run still lands in the inbox with a composed draft
    _, st = run_batch(s, stops_after_scoring, ["services-en-01"], False)
    item = st["items"][0]
    assert item["route"] == "negotiate" and item["pending"] == "approve_send"
    tr = [
        json.loads(x)
        for x in (s.runs_dir / item["run_id"] / "trace.jsonl").read_text().splitlines()
    ]
    assert len([e for e in tr if e["type"] == "nudge"]) == 2  # MAX_NUDGES, then it gives up
    inbox_item = next(i for i in W.inbox(s) if i["run_id"] == item["run_id"])
    draft = inbox_item["agent_text"]
    assert draft.startswith("Requested changes to Consulting Services Agreement")
    # cited by the contract's own numbering, not our clause id c3
    assert "\n1. Clause 3 (Termination): Either party may terminate" in draft
    # auto-send mode: the workflow sends the composed redline itself and says so
    s2 = settings(tmp_path / "auto")
    _, st2 = run_batch(s2, stops_after_scoring, ["services-en-01"], True)
    it2 = st2["items"][0]
    assert it2["pending"] is None and it2["decision"] == "send" and st2["needs_human"] == 0
    assert it2["decided_by"] == "workflow (auto-send policy)"
    rec = W.read_decision(s2, it2["run_id"])
    assert rec["by"] == "workflow (auto-send policy)" and "Requested changes" in rec["message"]


def test_unrouted_run_goes_to_a_lawyer_and_saves_no_time(tmp_path):
    s = settings(tmp_path)
    quits = [call_tools(("1", "load_contract", {"doc_id": "saas-en-01"})), text("I give up")]
    _, st = run_batch(s, lambda: list(quits), ["saas-en-01"], True)
    it = st["items"][0]
    assert it["status"] == "done" and it["route"] is None and it["pending"] == "unrouted"
    assert st["routes"]["unrouted"] == 1 and st["needs_human"] == 1
    assert st["estimate"]["hours_saved"] == 0.0  # a full manual review is still needed
    item = next(i for i in W.inbox(s) if i["run_id"] == it["run_id"])
    assert item["kind"] == "unrouted" and "lawyer must review" in item["reasons"][0]
    # red flags are still shown even though the model never scored
    assert {t["rule_id"] for t in item["tripwires"]} >= {"LIAB-CAP", "GOV-LAW"}
    assert any(t["status"] == "critical" for t in item["tripwires"])
    assert W.decide(s, it["run_id"], "reject")["kind"] == "unrouted"


def test_failed_model_call_is_a_clean_502_and_listed_with_its_contract(tmp_path):
    s = settings(tmp_path)

    def boom(_messages):
        raise RuntimeError("401 invalid x-api-key")

    client = TestClient(create_app(s, provider_factory=lambda: FakeProvider([boom])))
    r = client.post(
        "/run", json={"task": "Review contract nda-en-01 and route it.", "run_id": "ui-x-1"}
    )
    assert r.status_code == 502 and "model call failed" in r.json()["detail"]
    row = next(x for x in client.get("/runs").json() if x["run_id"] == "ui-x-1")
    assert row["status"] == "error" and row["doc_id"] == "nda-en-01"


def test_red_flag_excerpt_uses_the_contracts_own_words():
    from hack.challenge.contracts import normalize, normalize_with_map
    from hack.challenge.tripwires import scan

    s = Settings(provider="fake")
    for p in s.contracts_dir.glob("*.md"):
        contract = C.load_contract(s.contracts_dir, p.stem)
        for cl in contract.clauses:
            assert normalize_with_map(cl.text)[0] == normalize(
                cl.text
            )  # same text the patterns saw
        for t in scan(contract):
            assert t.match.replace(" …", "").split()[0] in contract.clause(t.clause_id).text
    ar = C.load_contract(s.contracts_dir, "ar-data-01")
    flag = next(t for t in scan(ar) if t.rule_id == "INTEREST")
    assert "فائدة" in flag.match  # original teh marbuta, not the normalised heh


def test_runs_without_a_contract_never_reach_the_inbox(tmp_path):
    s = settings(tmp_path)
    from hack.service import run_task

    r = run_task("hello", s, provider=FakeProvider([text("hi")]), run_id="ui-smoke-1")
    assert W.pending_kind(W.run_state(s, r.run_id), None) is None
    assert W.inbox(s) == []


def test_demo_replay_plays_a_recorded_run_without_any_model(tmp_path):
    import time as _t

    s = settings(tmp_path, demo_dir=Settings().demo_dir)
    client = TestClient(create_app(s))
    sc = client.get("/demo/scenes").json()
    assert {x["id"] for x in sc} >= {"approve-nda", "escalate-injection"}
    r = client.post("/demo/replay", json={"scene_id": "escalate-injection", "speed": 50}).json()
    rid = r["run_id"]
    assert rid.startswith("demo-escalate-injection-")
    for _ in range(200):
        rv = client.get(f"/runs/{rid}/review")
        if rv.status_code == 200 and rv.json()["status"] == "done":
            break
        _t.sleep(0.05)
    rv = rv.json()
    assert rv["route"] == "escalate" and rv["replay_of"] and rv["pending"] == "legal_review"
    assert any(i["run_id"] == rid for i in client.get("/inbox").json())
    assert client.post("/demo/replay", json={"scene_id": "nope"}).status_code == 404
    assert client.post("/demo/reset").json()["deleted"] == 1
    assert client.get(f"/runs/{rid}/review").status_code == 404


def test_learning_loop_turns_lawyer_decisions_into_a_playbook_amendment(tmp_path):
    from hack import learning as L
    from hack.challenge.tools import build_registry
    from hack.llm.base import ToolCall

    s = settings(tmp_path)
    client = TestClient(create_app(s))
    # two negotiate runs whose redlines legal decided NOT to send: the deviation is acceptable
    _, st = run_batch(s, lambda: scripted_review("services-en-01"), ["services-en-01"] * 2, False)
    for it in st["items"]:
        W.decide(s, it["run_id"], "discard", note="we accept supplier-only termination here")
    ins = client.get("/insights").json()
    assert ins["signals"]["TERM-CONV"] == {"accepted": 2, "upheld": 0}
    (p,) = ins["proposals"]
    assert (
        p["rule_id"] == "TERM-CONV"
        and "Client may terminate this Agreement only" in p["proposed_addition"]
    )
    assert (
        client.post(f"/insights/proposals/{p['id']}", json={"action": "maybe"}).status_code == 400
    )
    acc = client.post(
        f"/insights/proposals/{p['id']}", json={"action": "accept", "note": "ok"}
    ).json()
    assert acc["status"] == "accepted"
    assert client.get("/insights").json()["proposals"] == []  # not proposed again
    assert (
        client.post(f"/insights/proposals/{p['id']}", json={"action": "accept"}).status_code == 404
    )
    # the agent now sees the accepted wording in its playbook
    rules = json.loads(build_registry(s).execute(ToolCall("x", "list_playbook", {})).content)
    term = next(r for r in rules if r["id"] == "TERM-CONV")
    assert "Client may terminate this Agreement only" in term["accepted_by_legal"][0]
    assert L.amendments_by_rule(s)["TERM-CONV"]


def test_learning_loop_never_relaxes_critical_rules(tmp_path):
    from hack import learning as L

    s = settings(tmp_path)
    crit = [
        call_tools(("1", "load_contract", {"doc_id": "saas-en-01"})),
        call_tools(("2", "record_finding", {"doc_id": "saas-en-01", "clause_id": "c9", "rule_id": "LIAB-CAP", "quote": "Customer's aggregate liability", "issue": "uncapped", "redline": "cap", "confidence": 0.9})),
        text("stop"),
    ]  # fmt: skip
    _, st = run_batch(s, lambda: list(crit), ["saas-en-01"] * 2, False)
    for it in st["items"]:
        W.decide(s, it["run_id"], "approve_contract")
    ins = L.insights(s)
    assert ins["proposals"] == [] and ins["notes"][0]["rule_id"] == "LIAB-CAP"


def test_demo_replays_are_marked_in_the_inbox_and_never_teach_the_learning_loop(tmp_path):
    import time as _t

    from hack import learning as L

    s = settings(tmp_path, demo_dir=Settings().demo_dir)
    client = TestClient(create_app(s))
    rids = [
        client.post("/demo/replay", json={"scene_id": "negotiate-arabic", "speed": 50}).json()[
            "run_id"
        ]
        for _ in range(2)
    ]
    for rid in rids:
        for _ in range(200):
            if client.get(f"/runs/{rid}/review").json().get("status") == "done":
                break
            _t.sleep(0.05)
    items = {i["run_id"]: i for i in client.get("/inbox").json()}
    assert all(items[r]["replay_of"] for r in rids)
    for rid in rids:
        W.decide(s, rid, "discard")
    assert L.signals(s) == {} and L.insights(s)["proposals"] == []


def test_reasons_read_as_plain_text_even_for_old_traces():
    from hack.challenge.view import _clean_reason

    assert _clean_reason(
        "red flag(s) that always need a lawyer: ['LIAB-CAP@c8', 'GOV-LAW@c2']"
    ) == ("red flag(s) that always need a lawyer: LIAB-CAP@c8, GOV-LAW@c2")
    assert _clean_reason("Plain text, c7") == "Plain text, c7"


def test_index_is_revalidated_and_assets_are_immutable(tmp_path):
    s = settings(tmp_path)
    (s.web_dir / "assets").mkdir(parents=True)
    (s.web_dir / "index.html").write_text("<html>app</html>")
    (s.web_dir / "assets" / "index-abc.js").write_text("x")
    client = TestClient(create_app(s))
    assert client.get("/").headers["cache-control"] == "no-cache"
    assert "immutable" in client.get("/assets/index-abc.js").headers["cache-control"]


def test_counterparty_message_is_built_by_code_in_the_contract_language():
    from hack.challenge.messages import requested_changes

    f = [
        {"severity": "major", "clause_id": "c4", "clause_type": "payment", "redline": " Fee 1%. "},
        {"severity": "minor", "clause_id": "c5", "clause_type": "notices", "redline": "x"},
    ]
    ar = requested_changes("اتفاقية خدمات", f, {"c4": "المادة الرابعة: الدفع"})
    assert ar == "التعديلات المطلوبة على اتفاقية خدمات\n\n1. المادة الرابعة: Fee 1%."
    en = requested_changes("Services agreement", f, {"c4": "4. Payment. Net 30."})
    assert en == "Requested changes to Services agreement\n\n1. Clause 4 (Payment): Fee 1%."
    assert requested_changes("x", f[1:], {}) == ""  # minor only: nothing to ask for


def test_review_allowance_caps_single_runs_and_batches(tmp_path):
    """A shared demo sets HACK_MAX_REVIEWS so its API key cannot be used up: a batch that would
    pass the allowance starts nothing, and single runs stop at the cap. Replays do not count."""
    s = settings(tmp_path, max_reviews=3)
    client = TestClient(create_app(s, provider_factory=lambda: FakeProvider()))
    (s.runs_dir / "demo-x").mkdir(parents=True)
    (s.runs_dir / "demo-x" / "trace.jsonl").write_text("{}\n")  # a replay: free
    big = client.post(
        "/batches", json={"doc_ids": ["nda-en-01", "ar-data-01", "services-en-01", "ar-saas-01"]}
    )
    assert big.status_code == 429 and "3 are left" in big.json()["detail"]
    for i in range(3):
        r = client.post(
            "/run", json={"task": "Review contract nda-en-01 and route it.", "run_id": f"q-{i}"}
        )
        assert r.status_code == 200, r.text
    over = client.post(
        "/run", json={"task": "Review contract nda-en-01 and route it.", "run_id": "q-9"}
    )
    assert over.status_code == 429 and "0 are left" in over.json()["detail"]
    assert client.get("/healthz").json()["reviews"] == {"max": 3, "used": 3}
